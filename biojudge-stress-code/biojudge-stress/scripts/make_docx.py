"""Build the Word version of the paper from a lightweight markdown source + a finished run.

  python scripts/make_docx.py --src paper/paper_final.md --metrics runs/local/results/metrics.json \
      --tables paper/tables --out ../../BioJudge-Stress_paper_final.docx

Source syntax (one block per blank-line-separated paragraph):
  # Title / ## Section / ### Subsection
  AUTHOR: text                      centred author line
  - item                            bullet
  > text                            centred display equation
  {{TABLE:name|caption}}            table from <tables>/name.tex (written by biojudge.report)
  {{FIG:path|caption|width_in}}     figure (width optional; caption may contain '|')
  inline: **bold**, *italic*, x_{sub}, x^{sup}, [MacroName] -> value from metrics.json
"""
from __future__ import annotations

import argparse
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
TABLE_W = 6.5
FIRST_COL_W = {"tab_data": 1.9, "tab_bias_profile": 1.15, "tab_corrections": 1.9, "tab_regression": 2.4,
               "tab_attacks": 2.0, "tab_domain": 1.15}

TABLE_HEADERS = {
    "tab_data": [["Domain", "Sampled", "Kept", "Hard (%)", "Words", "Pad"]],
    "tab_bias_profile": [["Judge", "Acc", "κ", "Flip", "PPI", "VB", "SPB", "IC_{S×V}", "ASR_{combo}", "ASR_{any}"]],
    "tab_corrections": [["Configuration", "Acc", "Flip", "VB", "|SPB|", "ASR_{combo}", "ASR_{any}", "ΔAcc"]],
    "tab_regression": [["Term", "Coef. (pp)", "95% CI", "p"]],
    "tab_attacks": [["Edit", "Base", "+Debias", "+Norm", "+Deb+Norm", "Jury", "Full"]],
    "tab_domain": [["Judge", ("Acc (%)", 3), ("VB (pp)", 3), ("SPB (pp)", 3), ("ASR_{any} (%)", 3)],
                   [""] + ["L", "C", "P"] * 4],
}


# ----------------------------------------------------------------------------- macros
NON_MACROS = {"Question", "Answer"}
def clean_macro(v: str) -> str:
    v = v.replace("{,}", ",").replace("$<$", "<").replace("$p$", "p").replace("\\_", "_")
    v = v.replace("\\&", "&").replace("\\%", "%")
    v = re.sub(r"p<0\.001", "p < 0.001", v)
    v = re.sub(r"p=(\d)", r"p = \1", v)
    return v.replace("-", "−") if re.fullmatch(r"-\d+(\.\d+)?", v) else v


def fill_macros(text: str, macros: dict, missing: set) -> str:
    def rep(m):
        name = m.group(1)
        if name in macros:
            return clean_macro(macros[name])
        if name not in NON_MACROS:
            missing.add(name)
        return m.group(0)
    return re.sub(r"\[([A-Z][A-Za-z]+)\]", rep, text)


# ----------------------------------------------------------------------------- inline runs
_INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")


def add_runs(par, text: str, size=None, bold=False, italic=False):
    for tok in _INLINE.split(text):
        if not tok:
            continue
        b, i, sub, sup = bold, italic, False, False
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            tok, i = tok[1:-1], True
        elif tok.startswith("_{"):
            tok, sub = tok[2:-1], True
        elif tok.startswith("^{"):
            tok, sup = tok[2:-1], True
        # nested sub/superscripts inside bold/italic spans
        if (b != bold or i != italic) and _INLINE.search(tok):
            add_runs(par, tok, size, b, i)
            continue
        r = par.add_run(tok)
        r.bold, r.italic = b, i
        r.font.subscript, r.font.superscript = sub, sup
        if size:
            r.font.size = Pt(size)


# ----------------------------------------------------------------------------- tables
def _tex_cell(c: str):
    """Returns (text, bold, halfwidth, span)."""
    c = c.strip()
    span = 1
    m = re.match(r"\\multicolumn\{(\d+)\}\{c\}\{(.*)\}$", c)
    if m:
        span, c = int(m.group(1)), m.group(2)
    hw = None
    m = re.search(r"\{\\scriptsize\$\\pm\$([^}]*)\}", c)
    if m:
        hw = m.group(1)
        c = c[:m.start()] + c[m.end():]
    bold = False
    m = re.match(r"\\textbf\{(.*)\}$", c)
    if m:
        bold, c = True, m.group(1)
    c = (c.replace("$<$", "<").replace("$p$", "p").replace("$\\times$", "×").replace("\\_", "_")
          .replace("\\%", "%").replace("$|$", "|").replace("--", "–").replace("$", ""))
    c = re.sub(r"(?<![\w.])-(?=\d)", "−", c)
    return c, bold, hw, span


def parse_tex_rows(path: str):
    s = open(path, encoding="utf-8").read()
    body = s.split("\\midrule", 1)[1].split("\\bottomrule")[0]
    rows, rules = [], set()
    for line in body.split("\\\\"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("\\midrule"):
            rules.add(len(rows))
            line = line[len("\\midrule"):].strip()
        cells = [_tex_cell(c) for c in line.split("&")]
        rows.append(cells)
    return rows, rules


def _border(cell_or_tbl_pr, edges, sz=8):
    el = OxmlElement("w:tcBorders" if cell_or_tbl_pr.tag.endswith("tcPr") else "w:tblBorders")
    for e in edges:
        b = OxmlElement(f"w:{e}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), "000000")
        el.append(b)
    cell_or_tbl_pr.append(el)


def add_table(doc, name, tables_dir, caption, font_size=7.5):
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(cap, caption, size=8.5)
    rows, rules = parse_tex_rows(os.path.join(tables_dir, name + ".tex"))
    headers = TABLE_HEADERS[name]
    ncol = sum(h[1] if isinstance(h, tuple) else 1 for h in headers[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    _border(tblPr, ["top", "bottom"], sz=10)

    def fill(row_cells, spec, header=False, bottom_rule=False, top_rule=False):
        ci = 0
        for item in spec:
            if header:
                text, span = item if isinstance(item, tuple) else (item, 1)
                bold, hw = True, None
            else:
                text, bold, hw, span = item
            cell = row_cells[ci]
            if span > 1:
                cell = cell.merge(row_cells[ci + span - 1])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, text, size=font_size, bold=bold)
            if hw is not None:
                r = p.add_run(f" ±{hw}")
                r.font.size = Pt(font_size - 1.5)
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            edges = (["bottom"] if bottom_rule else []) + (["top"] if top_rule else [])
            if edges:
                for k in range(ci, ci + span):
                    _border(row_cells[k]._tc.get_or_add_tcPr(), edges, sz=6)
            ci += span

    first_w = FIRST_COL_W.get(name, 1.5)
    other_w = (TABLE_W - first_w) / (ncol - 1)
    for hi, h in enumerate(headers):
        fill(t.add_row().cells, h, header=True, bottom_rule=(hi == len(headers) - 1))
    for ri, r in enumerate(rows):
        fill(t.add_row().cells, r, top_rule=(ri in rules and ri > 0))
    t.autofit = False
    for row in t.rows:
        for k, c in enumerate(row.cells):
            c.width = Inches(first_w if k == 0 else other_w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_pipe_table(doc, lines, font_size=8, widths=(1.3, 5.2)):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines
            if not re.fullmatch(r"\|?[\s:|-]+\|?", l.strip())]
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _border(t._tbl.tblPr, ["top", "bottom"], sz=10)
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for k, text in enumerate(r):
            p = cells[k].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, text, size=font_size, bold=(ri == 0))
            cells[k].width = Inches(widths[k] if k < len(widths) else 1.0)
            if ri == 0:
                _border(cells[k]._tc.get_or_add_tcPr(), ["bottom"], sz=6)
    t.autofit = False
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ----------------------------------------------------------------------------- document
def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 17), ("Heading 1", 12), ("Heading 2", 10.5), ("Caption", 8.5)]:
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.bold = name != "Caption"
        s.font.italic = name == "Heading 2"
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    sec.top_margin = sec.bottom_margin = Inches(0.9)


def build(src, macros, tables_dir, out, base_dir):
    doc = Document()
    setup(doc)
    missing: set = set()
    text = fill_macros(open(src, encoding="utf-8").read(), macros, missing)
    blocks = [b.strip("\n") for b in re.split(r"\n\s*\n", text) if b.strip()]
    for b in blocks:
        if b.startswith("<!--"):
            continue
        line = " ".join(l.strip() for l in b.split("\n"))
        if b.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            add_runs(p, line[2:], size=17, bold=True)
        elif b.startswith("## "):
            add_runs(doc.add_paragraph(style="Heading 1"), line[3:], bold=True)
        elif b.startswith("### "):
            add_runs(doc.add_paragraph(style="Heading 2"), line[4:], italic=True)
        elif b.startswith("AUTHOR:"):
            for l in b.split("\n"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(1)
                add_runs(p, l.replace("AUTHOR:", "").strip(), size=10)
        elif b.startswith("- "):
            for item in re.split(r"\n(?=- )", b):
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, " ".join(l.strip() for l in item[2:].split("\n")))
        elif b.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, line[2:], italic=False)
        elif b.startswith("{{CAPTION:"):
            c = doc.add_paragraph(style="Caption")
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(c, line[10:-2], size=8.5)
        elif b.startswith("|"):
            add_pipe_table(doc, [l for l in b.split("\n") if l.strip()])
        elif b.startswith("{{TABLE:"):
            name, cap = line[8:-2].split("|", 1)
            add_table(doc, name, tables_dir, cap)
        elif b.startswith("{{FIG:"):
            # The caption may itself contain "|" (e.g. |SPB|), so the optional trailing width is
            # recognised by being numeric rather than by position.
            parts = line[6:-2].split("|")
            path, rest = parts[0], parts[1:]
            width = 6.0
            if len(rest) > 1:
                try:
                    width = float(rest[-1])
                    rest = rest[:-1]
                except ValueError:
                    pass
            cap = "|".join(rest)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(base_dir, path), width=Inches(width))
            c = doc.add_paragraph(style="Caption")
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(c, cap, size=8.5)
        elif b.startswith("REF:"):
            for l in re.split(r"\n(?=\[\d+\])", b[4:].strip()):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.3)
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, " ".join(x.strip() for x in l.split("\n")), size=8.5)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, line)
    doc.save(out)
    return missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--metrics", required=True)
    ap.add_argument("--tables", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--base", default=".", help="directory figure paths are relative to")
    a = ap.parse_args()
    macros = json.load(open(a.metrics, encoding="utf-8"))["macros"]
    missing = build(a.src, macros, a.tables, a.out, a.base)
    print(f"wrote {a.out}")
    if missing:
        print("UNFILLED TOKENS:", ", ".join(sorted(missing)))


if __name__ == "__main__":
    main()
